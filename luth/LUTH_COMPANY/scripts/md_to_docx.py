"""Convert the project's markdown specs into .docx files for client delivery.

Handles only the markdown subset used in docs/: headings, tables, bullet lists,
paragraphs and inline bold.

Usage:
    python scripts/md_to_docx.py docs/PRD_v1.2.md [more.md ...]
"""

import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

KOREAN_FONT = "맑은 고딕"
BOLD_PATTERN = re.compile(r"\*\*(.+?)\*\*")


def apply_korean_font(document):
    style = document.styles["Normal"]
    style.font.name = KOREAN_FONT
    style.font.size = Pt(10)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), KOREAN_FONT)


def write_inline(paragraph, text):
    """Split on **bold** markers and add runs with the right weight."""
    cursor = 0
    for match in BOLD_PATTERN.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor : match.start()])
        paragraph.add_run(match.group(1)).bold = True
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def split_row(line):
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def is_separator_row(line):
    return bool(re.fullmatch(r"\|[\s:|-]+\|", line.strip()))


def add_table(document, rows):
    header, *body = rows
    table = document.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, header):
        cell.text = ""
        write_inline(cell.paragraphs[0], text)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in body:
        cells = table.add_row().cells
        # Tolerate rows with fewer columns than the header.
        for cell, text in zip(cells, row):
            cell.text = ""
            write_inline(cell.paragraphs[0], text)
    document.add_paragraph()


def convert(md_path: Path, out_path: Path):
    document = Document()
    apply_korean_font(document)

    lines = md_path.read_text(encoding="utf-8").splitlines()
    pending_table = []

    def flush_table():
        if pending_table:
            add_table(document, pending_table)
            pending_table.clear()

    for line in lines:
        stripped = line.strip()

        if stripped.startswith("|"):
            if not is_separator_row(stripped):
                pending_table.append(split_row(stripped))
            continue
        flush_table()

        if not stripped or stripped == "---":
            continue

        if stripped.startswith("### "):
            document.add_heading(stripped[4:], level=3)
        elif stripped.startswith("## "):
            document.add_heading(stripped[3:], level=2)
        elif stripped.startswith("# "):
            document.add_heading(stripped[2:], level=1)
        elif stripped.startswith("- "):
            write_inline(document.add_paragraph(style="List Bullet"), stripped[2:])
        else:
            write_inline(document.add_paragraph(), stripped)

    flush_table()
    document.save(out_path)
    return out_path


def main():
    targets = sys.argv[1:]
    if not targets:
        print(__doc__)
        return 1

    for target in targets:
        md_path = Path(target)
        out_path = md_path.with_suffix(".docx")
        convert(md_path, out_path)
        print(f"created {out_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
